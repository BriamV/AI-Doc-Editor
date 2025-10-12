"""
Pytest fixtures for RAG services unit tests.

Provides mock objects and sample data for testing RAG pipeline services
without external dependencies (OpenAI, ChromaDB, file system).
"""

import pytest
from pathlib import Path
from typing import List, Dict, Any
from unittest.mock import Mock, MagicMock
import tempfile
import io


@pytest.fixture
def temp_test_dir():
    """Create a temporary directory for test files."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def sample_pdf_path(temp_test_dir: Path) -> Path:
    """
    Create a minimal valid PDF file for testing.

    Note: This is a minimal PDF that PyPDF2 can parse.
    """
    pdf_content = b"""%PDF-1.4
1 0 obj
<<
/Type /Catalog
/Pages 2 0 R
>>
endobj
2 0 obj
<<
/Type /Pages
/Kids [3 0 R]
/Count 1
>>
endobj
3 0 obj
<<
/Type /Page
/Parent 2 0 R
/MediaBox [0 0 612 792]
/Contents 4 0 R
/Resources <<
/Font <<
/F1 <<
/Type /Font
/Subtype /Type1
/BaseFont /Helvetica
>>
>>
>>
>>
endobj
4 0 obj
<<
/Length 44
>>
stream
BT
/F1 12 Tf
100 700 Td
(Sample PDF text) Tj
ET
endstream
endobj
xref
0 5
0000000000 65535 f
0000000009 00000 n
0000000058 00000 n
0000000115 00000 n
0000000317 00000 n
trailer
<<
/Size 5
/Root 1 0 R
>>
startxref
410
%%EOF
"""
    pdf_path = temp_test_dir / "sample.pdf"
    pdf_path.write_bytes(pdf_content)
    return pdf_path


@pytest.fixture
def sample_docx_path(temp_test_dir: Path) -> Path:
    """
    Create a minimal valid DOCX file for testing.

    Note: Creates an actual DOCX using python-docx library.
    """
    from docx import Document as DocxDocument

    docx_path = temp_test_dir / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("Sample DOCX paragraph 1")
    doc.add_paragraph("Sample DOCX paragraph 2")

    # Add a table for testing table extraction
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"
    table.rows[1].cells[0].text = "Cell 1"
    table.rows[1].cells[1].text = "Cell 2"

    doc.save(str(docx_path))
    return docx_path


@pytest.fixture
def sample_markdown_path(temp_test_dir: Path) -> Path:
    """Create a sample Markdown file for testing."""
    md_content = """# Sample Markdown

This is a paragraph with **bold** and *italic* text.

## Section 2

- List item 1
- List item 2
- List item 3

Another paragraph with some content.
"""
    md_path = temp_test_dir / "sample.md"
    md_path.write_text(md_content, encoding="utf-8")
    return md_path


@pytest.fixture
def empty_file_path(temp_test_dir: Path) -> Path:
    """Create an empty file for testing."""
    empty_path = temp_test_dir / "empty.txt"
    empty_path.write_text("", encoding="utf-8")
    return empty_path


@pytest.fixture
def sample_text_short() -> str:
    """Short text sample (< chunk_size) for testing."""
    return "This is a short text sample that fits in one chunk."


@pytest.fixture
def sample_text_long() -> str:
    """Long text sample (> chunk_size) for testing chunking."""
    # Create text longer than default chunk_size (1000 chars)
    paragraphs = [
        "This is the first paragraph. " * 20,
        "This is the second paragraph. " * 20,
        "This is the third paragraph. " * 20,
        "This is the fourth paragraph. " * 20,
    ]
    return "\n\n".join(paragraphs)


@pytest.fixture
def sample_chunks() -> List[str]:
    """Sample text chunks for embedding tests."""
    return [
        "This is the first chunk of text for testing embeddings.",
        "This is the second chunk with different content.",
        "A third chunk to test batch processing.",
    ]


@pytest.fixture
def mock_openai_client():
    """
    Mock OpenAI client for embedding service tests.

    Returns a properly structured mock that simulates OpenAI API responses.
    """
    mock_client = Mock()

    # Mock embeddings.create response
    mock_embedding_data = Mock()
    mock_embedding_data.embedding = [0.1] * 1536  # text-embedding-3-small dimension

    mock_response = Mock()
    mock_response.data = [mock_embedding_data]
    mock_response.usage.total_tokens = 100

    mock_client.embeddings.create.return_value = mock_response

    return mock_client


@pytest.fixture
def mock_openai_client_rate_limit():
    """Mock OpenAI client that simulates rate limiting."""
    from openai import RateLimitError

    mock_client = Mock()
    mock_client.embeddings.create.side_effect = RateLimitError("Rate limit exceeded")

    return mock_client


@pytest.fixture
def mock_openai_client_connection_error():
    """Mock OpenAI client that simulates connection errors."""
    from openai import APIConnectionError

    mock_client = Mock()
    mock_client.embeddings.create.side_effect = APIConnectionError("Connection failed")

    return mock_client


@pytest.fixture
def mock_chroma_client():
    """
    Mock ChromaDB client for vector store tests.

    Returns a properly structured mock for ChromaDB operations.
    """
    mock_client = Mock()
    mock_collection = Mock()

    # Mock collection methods
    mock_collection.add = Mock(return_value=None)
    mock_collection.query = Mock(return_value={
        "documents": [["Sample document text"]],
        "distances": [[0.1]],
        "metadatas": [[{"document_id": "doc-1", "user_id": "user-1"}]],
        "ids": [["chunk_0"]],
    })
    mock_collection.delete = Mock(return_value=None)
    mock_collection.count = Mock(return_value=10)
    mock_collection.update = Mock(return_value=None)
    mock_collection.name = "test_collection"

    # Mock client methods
    mock_client.get_collection = Mock(return_value=mock_collection)
    mock_client.create_collection = Mock(return_value=mock_collection)
    mock_client.delete_collection = Mock(return_value=None)
    mock_client.list_collections = Mock(return_value=[mock_collection])

    return mock_client


@pytest.fixture
def sample_embeddings() -> List[List[float]]:
    """Sample embedding vectors for testing."""
    return [
        [0.1] * 1536,  # Embedding for chunk 1
        [0.2] * 1536,  # Embedding for chunk 2
        [0.3] * 1536,  # Embedding for chunk 3
    ]


@pytest.fixture
def sample_metadata() -> List[Dict[str, Any]]:
    """Sample metadata for document chunks."""
    return [
        {
            "document_id": "doc-123",
            "user_id": "user-456",
            "chunk_index": 0,
            "total_chunks": 3,
            "source_file": "sample.pdf",
        },
        {
            "document_id": "doc-123",
            "user_id": "user-456",
            "chunk_index": 1,
            "total_chunks": 3,
            "source_file": "sample.pdf",
        },
        {
            "document_id": "doc-123",
            "user_id": "user-456",
            "chunk_index": 2,
            "total_chunks": 3,
            "source_file": "sample.pdf",
        },
    ]


@pytest.fixture
def mock_db_session():
    """Mock database session for RAG processing service tests."""
    mock_session = MagicMock()
    return mock_session


@pytest.fixture
def mock_document_service():
    """Mock DocumentService for RAG processing tests."""
    mock_service = Mock()

    # Mock update_document_status as async
    async def mock_update_status(*args, **kwargs):
        return {"status": "success"}

    mock_service.update_document_status = Mock(side_effect=mock_update_status)

    return mock_service
